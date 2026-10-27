"""
Excel pipeline engine: messy text file -> verified .xlsx, with a reconciliation ledger.
Plain module (not a tool file). tools/excel.py exposes it to the agent.

Stages (deterministic Python; the LLM only proposes the rules):
  inspect -> immutable snapshot + SHA-256 -> profile -> [LLM rules] -> validate on sample
  -> streaming parse -> canonical records on disk -> explicit cleaning (every change logged)
  -> split valid / review / structural -> write xlsx -> reopen -> verify 3 levels -> QA report
"""
import codecs
import csv
import datetime as dt
import hashlib
import io
import json
import os
import random
import re
import shutil
import time
from collections import Counter, deque

import config
import session

csv.field_size_limit(2 ** 30)

MAX_ROWS = 1_048_576          # Excel sheet row limit (header included)
MAX_CHARS = 32_767            # Excel cell text limit
MAX_INT = 999_999_999_999_999  # 15 digits: Excel number precision
MIN_DATE = dt.date(1900, 3, 1)
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")   # cannot be stored in xlsx
_validated = {}               # (rules hash, source sha) -> dry-run summary (gate for build)


# =========================================================== basics
def _root():
    return os.path.abspath(config.WORKDIR)


def resolve_source(path):
    p = os.path.expanduser(path)
    if not os.path.isabs(p):
        p = os.path.join(_root(), p)
    p = os.path.abspath(p)
    if not os.path.isfile(p):
        raise FileNotFoundError(f"source file not found: {p}")
    return p


def resolve_output(path, source):
    if not path:
        path = os.path.splitext(os.path.basename(source))[0] + "_clean.xlsx"
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"
    return session.safe_path(path)


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _decodes(path, enc):
    dec = codecs.getincrementaldecoder(enc)("strict")
    try:
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                dec.decode(chunk)
        dec.decode(b"", final=True)
        return True
    except UnicodeDecodeError:
        return False


def detect_encoding(path):
    with open(path, "rb") as f:
        head = f.read(4)
    for bom, enc in ((b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe\x00\x00", "utf-32"),
                     (b"\x00\x00\xfe\xff", "utf-32"), (b"\xff\xfe", "utf-16"),
                     (b"\xfe\xff", "utf-16")):
        if head.startswith(bom):
            return enc, "byte-order mark"
    if _decodes(path, "utf-8"):
        return "utf-8", "valid utf-8"
    if _decodes(path, "cp1252"):
        return "cp1252", "NOT valid utf-8; decodes as cp1252 (check that accents look right)"
    return "latin-1", "fallback: any byte decodes (check that accents look right)"


def take_snapshot(src):
    """Read-only, content-addressed copy of the source. Returns (snapshot_path, sha256)."""
    sha = sha256_file(src)
    d = os.path.join(_root(), ".excel_snapshots", sha[:16])
    os.makedirs(d, exist_ok=True)
    snap = os.path.join(d, os.path.basename(src))
    if not (os.path.exists(snap) and sha256_file(snap) == sha):
        if os.path.exists(snap):
            os.chmod(snap, 0o666)
            os.remove(snap)
        tmp = snap + ".tmp"
        shutil.copyfile(src, tmp)
        if sha256_file(tmp) != sha:
            os.remove(tmp)
            raise RuntimeError("snapshot differs from source (file changing while it was copied?)")
        os.replace(tmp, snap)
        os.chmod(snap, 0o444)
    return snap, sha


def _short(s, n=40):
    s = repr(s)
    return s if len(s) <= n else s[:n - 1] + "…"


# =========================================================== cell encoding (type aware)
def enc_cell(v):
    """Type-aware encoding: "0012" (text) != 12 (int); None != "" is moot in Excel (blank)."""
    if v is None:
        return ["n"]
    if isinstance(v, bool):
        return ["b", v]
    if isinstance(v, int):
        return ["i", str(v)]
    if isinstance(v, float):
        return ["f", "%.15g" % v]            # Excel keeps 15 significant digits
    if isinstance(v, dt.datetime):
        return ["t", v.isoformat()]
    if isinstance(v, dt.date):
        return ["d", v.isoformat()]
    if isinstance(v, str):
        return ["s", v]
    raise TypeError(f"unsupported cell type {type(v)}")


def dec_cell(e):
    t = e[0]
    if t == "n":
        return None
    if t == "s":
        return e[1]
    if t == "i":
        return int(e[1])
    if t == "f":
        return float(e[1])
    if t == "b":
        return bool(e[1])
    if t == "d":
        return dt.date.fromisoformat(e[1])
    return dt.datetime.fromisoformat(e[1])


def _row_hash(ref, cells):
    blob = json.dumps([ref, cells], separators=(",", ":"), ensure_ascii=True)
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()


def _back(v, ctype):
    """Normalise what openpyxl reads back so it can be compared with what we meant to write."""
    if v is None:
        return None
    if ctype == "date" and isinstance(v, dt.datetime) and v.time() == dt.time(0):
        return v.date()
    if ctype == "datetime" and isinstance(v, dt.datetime):
        if v.microsecond >= 500_000:
            v = v + dt.timedelta(seconds=1)
        return v.replace(microsecond=0)
    if ctype == "decimal" and isinstance(v, int) and not isinstance(v, bool):
        return float(v)
    return v


# =========================================================== rules
_FORMATS = {"delimited", "fixed_width", "regex", "blocks", "jsonl"}
_TYPES = {"text", "integer", "decimal", "date", "datetime", "boolean"}
_OPS = {"trim": lambda s: s.strip(), "collapse_spaces": lambda s: re.sub(r"\s+", " ", s),
        "lower": lambda s: s.lower(), "upper": lambda s: s.upper(), "title": lambda s: s.title()}
_DEFAULTS = {
    "encoding": None, "format": None, "delimiter": ",", "quotechar": '"', "header_lines": 0,
    "skip_blank": True, "skip_regex": [], "fixed_width": None, "regex": None,
    "block_start": None, "block_kv_regex": r"^\s*([^:]+?)\s*:\s?(.*)$", "max_record_lines": 50,
    "columns": [], "ignored_sources": [], "include_source_ref": True, "sheet_name": "Data",
    "log_transformations": "changes"}
_COL_KEYS = {"name", "source", "type", "clean", "replace", "null_values", "required", "formats",
             "thousands", "decimal_comma", "true_values", "false_values"}

RULES_HELP = """RULES = one JSON object. Everything is explicit; nothing is dropped or guessed silently.
format: delimited | fixed_width | regex | blocks | jsonl
encoding: take it from excel_inspect (e.g. "utf-8")
delimited: delimiter (",", "\\t", ";", "|"), quotechar ('"' or null), header_lines (leading header lines, kept as structural), max_record_lines (multi-line quoted fields, default 50)
fixed_width: "fixed_width": [[start,end],...] character slices (end may be null = to end of line)
regex: "regex" with named groups; the WHOLE line must match, otherwise the line goes to review
blocks: a record = "Key: value" lines between blank lines; optional "block_start" regex begins a new record; "block_kv_regex" needs 2 groups (key, value)
jsonl: one JSON object per line
common: skip_blank (default true) · skip_regex [patterns]: lines classified as structural (comments, page headers), kept in the audit log · ignored_sources [..]: fields you deliberately do not export (non-empty values are logged) · include_source_ref (default true: first column = source line reference) · sheet_name · log_transformations: "changes" (default) | "full" | "summary"
columns: [{"name","source","type", ...}]
  source: field index (delimited, fixed_width) or group/key name (regex, blocks, jsonl)
  type: text (default; keeps IDs, zip codes, leading zeros) | integer | decimal | date | datetime | boolean
  clean: ["trim","collapse_spaces","lower","upper","title"]  applied in order, every change is logged
  replace: [{"pattern": regex, "with": text}]
  null_values: ["N/A","-"]  exact values turned into blank cells (logged)
  required: true -> a blank value sends the record to review
  integer/decimal: "thousands": ",", "decimal_comma": true
  date/datetime: "formats": ["%d/%m/%Y", ...] (strptime). Unmatched or ambiguous -> review. Years before 1900-03 and fractional seconds -> review
  boolean: "true_values", "false_values" (defaults: true/yes/y/1 and false/no/n/0)
Records that cannot be parsed cleanly, have unmapped non-empty fields, or fail a type conversion go to the Review sheet with their raw text. Fix the rules (or ask the user) instead of accepting a large review queue."""


class Rules:
    def __init__(self, raw):
        if isinstance(raw, str):
            try:
                raw = json.loads(raw)
            except ValueError as e:
                raise ValueError(f"rules is not valid JSON: {e}")
        if not isinstance(raw, dict):
            raise ValueError("rules must be a JSON object")
        unknown = set(raw) - set(_DEFAULTS)
        if unknown:
            raise ValueError(f"unknown rule keys {sorted(unknown)}; allowed: {sorted(_DEFAULTS)}")
        d = json.loads(json.dumps(_DEFAULTS))
        d.update(raw)
        self.fmt = d["format"]
        if self.fmt not in _FORMATS:
            raise ValueError(f"format must be one of {sorted(_FORMATS)}")
        if d["encoding"]:
            try:
                codecs.lookup(d["encoding"])
            except LookupError:
                raise ValueError(f"unknown encoding {d['encoding']!r}")
        delim = "\t" if d["delimiter"] in ("tab", "\\t") else d["delimiter"]
        if self.fmt == "delimited" and (not isinstance(delim, str) or len(delim) != 1):
            raise ValueError("delimiter must be one character")
        d["delimiter"] = delim
        if d["quotechar"] in ("", "none", None):
            d["quotechar"] = None
        elif not isinstance(d["quotechar"], str) or len(d["quotechar"]) != 1:
            raise ValueError("quotechar must be one character or null")
        for k in ("header_lines", "max_record_lines"):
            if not isinstance(d[k], int) or d[k] < 0:
                raise ValueError(f"{k} must be an integer >= 0")
        if d["log_transformations"] not in ("summary", "changes", "full"):
            raise ValueError('log_transformations must be "summary", "changes" or "full"')
        sheet = re.sub(r"[\[\]:*?/\\]", "_", str(d["sheet_name"]))[:31] or "Data"
        if sheet in ("Review", "Summary"):
            raise ValueError('sheet_name cannot be "Review" or "Summary"')
        d["sheet_name"] = sheet
        try:
            self.skip_rx = [re.compile(p) for p in d["skip_regex"]]
            self.rx = re.compile(d["regex"]) if d["regex"] else None
            self.block_start = re.compile(d["block_start"]) if d["block_start"] else None
            self.block_kv = re.compile(d["block_kv_regex"])
        except (re.error, TypeError) as e:
            raise ValueError(f"bad regex in rules: {e}")
        if self.fmt == "regex" and not (self.rx and self.rx.groupindex):
            raise ValueError('format "regex" needs "regex" with named groups (?P<name>...)')
        if self.fmt == "blocks" and self.block_kv.groups != 2:
            raise ValueError("block_kv_regex needs exactly 2 groups (key, value)")
        self.fw, self.fw_gaps = [], []
        if self.fmt == "fixed_width":
            fw = d["fixed_width"]
            if not fw or not all(isinstance(s, list) and len(s) == 2 and isinstance(s[0], int)
                                 and s[0] >= 0 and (s[1] is None or (isinstance(s[1], int) and s[1] > s[0]))
                                 for s in fw):
                raise ValueError('fixed_width must be [[start,end],...] with end > start (end may be null)')
            self.fw = [tuple(s) for s in fw]
            pos = 0
            for a, b in sorted(self.fw, key=lambda s: s[0]):
                if a > pos:
                    self.fw_gaps.append((pos, a))
                pos = max(pos, 10 ** 9 if b is None else b)
            if pos < 10 ** 9:
                self.fw_gaps.append((pos, None))
        by_index = self.fmt in ("delimited", "fixed_width")
        if not isinstance(d["columns"], list) or not d["columns"]:
            raise ValueError("columns must be a non-empty list")
        self.cols, names = [], set()
        for i, c in enumerate(d["columns"]):
            if not isinstance(c, dict) or set(c) - _COL_KEYS:
                raise ValueError(f"column {i}: unknown keys {sorted(set(c) - _COL_KEYS) if isinstance(c, dict) else c}")
            name = str(c.get("name", "")).strip()
            if not name or name in names or (d["include_source_ref"] and name == "source_ref"):
                raise ValueError(f"column {i}: name missing, duplicate or reserved ({name!r})")
            names.add(name)
            src = c.get("source")
            if by_index and (not isinstance(src, int) or isinstance(src, bool) or src < 0):
                raise ValueError(f"column {name}: source must be a field index (integer) for {self.fmt}")
            if not by_index and not isinstance(src, str):
                raise ValueError(f"column {name}: source must be a group/key name (string) for {self.fmt}")
            typ = c.get("type", "text")
            if typ not in _TYPES:
                raise ValueError(f"column {name}: type must be one of {sorted(_TYPES)}")
            clean = c.get("clean", [])
            if not isinstance(clean, list) or any(o not in _OPS for o in clean):
                raise ValueError(f"column {name}: clean ops must be from {sorted(_OPS)}")
            reps = []
            for r in c.get("replace", []):
                try:
                    reps.append((re.compile(r["pattern"]), r.get("with", "")))
                except (re.error, KeyError, TypeError) as e:
                    raise ValueError(f"column {name}: bad replace rule: {e}")
            if typ in ("date", "datetime") and not (c.get("formats") and isinstance(c["formats"], list)):
                raise ValueError(f"column {name}: {typ} needs a non-empty 'formats' list")
            self.cols.append({
                "name": name, "source": src, "type": typ, "clean": clean, "replace": reps,
                "null_values": set(c.get("null_values", [])), "required": bool(c.get("required", False)),
                "formats": c.get("formats", []), "thousands": c.get("thousands", "") or "",
                "decimal_comma": bool(c.get("decimal_comma", False)),
                "true": {str(x).lower() for x in c.get("true_values", ["true", "yes", "y", "1"])},
                "false": {str(x).lower() for x in c.get("false_values", ["false", "no", "n", "0"])}})
        self.ignored = set(d["ignored_sources"])
        self.mapped = {c["source"] for c in self.cols}
        self.delim, self.quote = d["delimiter"], d["quotechar"]
        self.header_lines, self.skip_blank = d["header_lines"], d["skip_blank"]
        self.max_rec, self.include_ref = d["max_record_lines"], d["include_source_ref"]
        self.sheet, self.tlevel = sheet, d["log_transformations"]
        self.enc = d["encoding"]
        self.d = d
        self.hash = hashlib.sha256(json.dumps(d, sort_keys=True).encode()).hexdigest()


# =========================================================== line reading + record assembly
class _Lines:
    def __init__(self, path, enc):
        self.f = open(path, "r", encoding=enc, newline="")
        self.it = enumerate(self.f, 1)
        self.back = deque()
        self.last = 0

    def next(self):
        if self.back:
            item = self.back.popleft()
        else:
            try:
                item = next(self.it)
            except StopIteration:
                return None
        self.last = max(self.last, item[0])
        return item

    def push(self, items):
        self.back.extendleft(reversed(items))

    def close(self):
        self.f.close()


def _eol(s):
    if s.endswith("\r\n"):
        return s[:-2]
    if s.endswith("\n") or s.endswith("\r"):
        return s[:-1]
    return s


def _structural(text, R, st):
    if R.skip_blank and not text.strip():
        return "blank_line"
    for rx in R.skip_rx:
        if rx.search(text):
            return "skip_regex"
    if st["header"] > 0:
        st["header"] -= 1
        return "header"
    return None


def _unit(start, end, kind, raw, reason="", fields=None):
    return {"start": start, "end": end, "kind": kind, "raw": raw, "reason": reason, "fields": fields}


def _read_delimited(L, n, raw, R):
    buf, used = raw, [(n, raw)]
    while True:
        err = None
        try:
            rows = list(csv.reader(io.StringIO(buf, newline=""), delimiter=R.delim,
                                   quotechar=R.quote or '"',
                                   quoting=csv.QUOTE_MINIMAL if R.quote else csv.QUOTE_NONE,
                                   strict=True))
            if len(rows) == 1:
                return _unit(n, used[-1][0], "data", "\n".join(_eol(r) for _, r in used), fields=rows[0])
            err = f"csv produced {len(rows)} rows from one record"
        except csv.Error as e:
            err = str(e)
            if "end of data" in err and len(used) < R.max_rec:
                nxt = L.next()
                if nxt is not None:
                    used.append(nxt)
                    buf += nxt[1]
                    continue
        L.push(used[1:])      # give the following lines back: only the first line is the bad record
        return _unit(n, n, "review", _eol(raw), f"csv_parse_error: {err}")


def _dup_hook(pairs):
    keys = [k for k, _ in pairs]
    if len(keys) != len(set(keys)):
        raise ValueError("duplicate key in JSON object")
    return dict(pairs)


def _units_lines(R, L):
    st = {"header": R.header_lines}
    while True:
        item = L.next()
        if item is None:
            return
        n, raw = item
        text = _eol(raw)
        reason = _structural(text, R, st)
        if reason:
            yield _unit(n, n, "structural", text, reason)
            continue
        if R.fmt == "delimited":
            yield _read_delimited(L, n, raw, R)
        elif R.fmt == "fixed_width":
            out = [i for i in R.fw_gaps if text[i[0]:i[1]].strip()]
            if out:
                yield _unit(n, n, "review", text, f"text_outside_columns at {out[0][0]}")
            else:
                yield _unit(n, n, "data", text, fields=[text[a:b] for a, b in R.fw])
        elif R.fmt == "regex":
            m = R.rx.fullmatch(text)
            if m:
                yield _unit(n, n, "data", text, fields=m.groupdict())
            else:
                yield _unit(n, n, "review", text, "regex_no_match")
        else:  # jsonl
            try:
                obj = json.loads(text, object_pairs_hook=_dup_hook)
            except ValueError as e:
                yield _unit(n, n, "review", text, f"invalid_json: {e}")
                continue
            if isinstance(obj, dict):
                yield _unit(n, n, "data", text, fields=obj)
            else:
                yield _unit(n, n, "review", text, "json_not_an_object")


def _make_block(cur, R):
    fields, bad = {}, None
    for n, text in cur:
        m = R.block_kv.match(text)
        if not m:
            bad = f"line {n} is not a key/value line"
            break
        k = m.group(1).strip()
        if k in fields:
            bad = f"duplicate key {k!r}"
            break
        fields[k] = m.group(2)
    raw = "\n".join(t for _, t in cur)
    if bad:
        return _unit(cur[0][0], cur[-1][0], "review", raw, f"block_error: {bad}")
    return _unit(cur[0][0], cur[-1][0], "data", raw, fields=fields)


def _units_blocks(R, L):
    cur = []
    while True:
        item = L.next()
        end = item is None
        blank = skip = starts = False
        if not end:
            n, raw = item
            text = _eol(raw)
            blank = not text.strip()
            skip = (not blank) and any(rx.search(text) for rx in R.skip_rx)
            starts = (not blank) and bool(R.block_start and R.block_start.search(text))
        if cur and (end or blank or skip or starts):
            yield _make_block(cur, R)
            cur = []
        if end:
            return
        if blank:
            yield _unit(n, n, "structural", text, "blank_line")
        elif skip:
            yield _unit(n, n, "structural", text, "skip_regex")
        else:
            cur.append((n, text))


def _iter_units(R, L):
    return _units_blocks(R, L) if R.fmt == "blocks" else _units_lines(R, L)


# =========================================================== cell conversion + cleaning
def _to_int(s, c):
    t = s.strip().replace(c["thousands"], "") if c["thousands"] else s.strip()
    if not re.fullmatch(r"[+-]?\d+", t):
        raise ValueError(f"{_short(s)} is not an integer")
    digits = t.lstrip("+-")
    if len(digits) > 1 and digits[0] == "0":
        raise ValueError(f"{_short(s)} has leading zeros; declare the column as text to keep them")
    v = int(t)
    if abs(v) > MAX_INT:
        raise ValueError(f"{_short(s)} exceeds 15 digits (Excel precision); declare the column as text")
    return v


def _to_dec(s, c):
    t = s.strip()
    if c["thousands"]:
        t = t.replace(c["thousands"], "")
    if c["decimal_comma"]:
        t = t.replace(",", ".")
    if not re.fullmatch(r"[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?", t):
        raise ValueError(f"{_short(s)} is not a decimal number")
    mant = re.split("[eE]", t)[0].lstrip("+-").replace(".", "").lstrip("0")
    if len(mant) > 15:
        raise ValueError(f"{_short(s)} has more than 15 significant digits; declare the column as text")
    v = float(t)
    if v != v or v in (float("inf"), float("-inf")):
        raise ValueError(f"{_short(s)} is not a finite number")
    return v


def _to_when(s, c, with_time):
    t = s.strip()
    found = set()
    for fmt in c["formats"]:
        try:
            found.add(dt.datetime.strptime(t, fmt))
        except ValueError:
            pass
    if not found:
        raise ValueError(f"{_short(s)} matches none of the declared formats")
    if len(found) > 1:
        raise ValueError(f"{_short(s)} is ambiguous across the declared formats")
    d = found.pop()
    if d.date() < MIN_DATE:
        raise ValueError(f"{_short(s)} is before 1900-03-01 (not safely representable in Excel); use text")
    if with_time:
        if d.microsecond:
            raise ValueError(f"{_short(s)} has fractional seconds; use text")
        return d
    if d.time() != dt.time(0):
        raise ValueError(f"{_short(s)} has a time part; use type datetime")
    return d.date()


def _to_bool(s, c):
    t = s.strip().lower()
    if t in c["true"]:
        return True
    if t in c["false"]:
        return False
    raise ValueError(f"{_short(s)} is not a declared true/false value")


_CONVERT = {"integer": _to_int, "decimal": _to_dec, "boolean": _to_bool,
            "date": lambda s, c: _to_when(s, c, False), "datetime": lambda s, c: _to_when(s, c, True)}


class _TL:
    """Transformation log: counts everything, writes per-cell lines according to the level."""

    def __init__(self, level, fh):
        self.level, self.fh, self.counts = level, fh, Counter()

    def count(self, op, n=1):
        self.counts[op] += n

    def __call__(self, ref, col, op, before, after):
        self.counts[op] += 1
        if self.fh and self.level != "summary":
            self.fh.write(json.dumps({"ref": ref, "col": col, "op": op, "before": before,
                                      "after": after if (after is None or isinstance(after, str)) else str(after)},
                                     ensure_ascii=True) + "\n")


def _scalar_text(v, tl=None, ref=None, col=None):
    if v is None or isinstance(v, str):
        return v
    if isinstance(v, (dict, list)):
        s = json.dumps(v, ensure_ascii=False, separators=(",", ":"))
        if tl:
            tl(ref, col, "json_to_text", "<nested JSON>", s)
        return s
    return json.dumps(v)


def _cell(raw, c, ref, tl):
    val, name = raw, c["name"]
    if val is not None:
        for op in c["clean"]:
            new = _OPS[op](val)
            if new != val:
                tl(ref, name, op, val, new)
                val = new
        for rx, rep in c["replace"]:
            new = rx.sub(rep, val)
            if new != val:
                tl(ref, name, "replace", val, new)
                val = new
        if c["null_values"] and val in c["null_values"]:
            tl(ref, name, "null_value", val, None)
            val = None
    if val == "":
        tl.count("empty_to_blank")
        val = None
    if val is None:
        if c["required"]:
            raise ValueError("required value is empty")
        return None
    if c["type"] == "text":
        if "\r" in val:
            val = val.replace("\r\n", "\n").replace("\r", "\n")
            tl.count("newline_normalized")
        if _ILLEGAL.search(val):
            raise ValueError("contains control characters that Excel cannot store")
        if len(val) > MAX_CHARS:
            raise ValueError(f"text longer than Excel's {MAX_CHARS} character cell limit")
        return val
    out = _CONVERT[c["type"]](val, c)
    if tl.level == "full":
        tl(ref, name, "convert:" + c["type"], val, out)
    else:
        tl.count("convert:" + c["type"])
    return out


def _build_row(u, R, ref, tl):
    fields, reasons, cells = u["fields"], [], []
    is_list = isinstance(fields, list)
    items = enumerate(fields) if is_list else fields.items()
    for k, v in items:
        txt = _scalar_text(v)
        if txt is None or not txt.strip():
            continue
        if k in R.ignored:
            tl(ref, f"[{k}]", "ignored_source", txt, None)
        elif k not in R.mapped:
            reasons.append(f"unmapped_{'field' if is_list else 'key'}_{k}: {_short(txt, 30)}")
    for c in R.cols:
        src = c["source"]
        if is_list:
            if src >= len(fields):
                reasons.append(f"missing_field_{src}")
                cells.append(None)
                continue
            raw = fields[src]
        else:
            if src not in fields:
                tl.count("missing_key_to_blank")
            raw = fields.get(src)
        try:
            cells.append(_cell(_scalar_text(raw, tl, ref, c["name"]), c, ref, tl))
        except ValueError as e:
            reasons.append(f"{c['name']}: {e}")
            cells.append(None)
    if not reasons and all(v is None for v in cells):
        reasons.append("all_values_empty")
    return cells, reasons


def _sheet_text(s):
    esc = _ILLEGAL.sub(lambda m: "\\x%02x" % ord(m.group()), s)
    if len(esc) > MAX_CHARS:
        esc = esc[:MAX_CHARS - 70] + " …[TRUNCATED in sheet; full text in audit/review_raw.jsonl]"
    return esc


REVIEW_HEADER = ["source_ref", "line_start", "line_end", "reason", "raw_text"]
REVIEW_TYPES = ["text", "integer", "integer", "text", "text"]


# =========================================================== the pipeline pass
def run_pipeline(R, snap, enc, limit=0, out_dir=None):
    """One streaming pass. Writes canonical records + audit files when out_dir is given."""
    led = {"units": 0, "data": 0, "review": 0, "structural": Counter(), "review_reasons": Counter(),
           "coverage_ok": True, "coverage_note": "", "complete": False, "lines": 0,
           "examples": {}, "sample_rows": [], "trans": None}
    fh = {}
    if out_dir:
        for name in ("expected_rows.jsonl", "review_rows.jsonl", "review_raw.jsonl",
                     "structural.jsonl", "transformations.jsonl"):
            fh[name] = open(os.path.join(out_dir, name), "w", encoding="utf-8", newline="\n")
        fh["row_map"] = open(os.path.join(out_dir, "row_map.csv"), "w", encoding="utf-8", newline="")
        fh["review_map"] = open(os.path.join(out_dir, "review_map.csv"), "w", encoding="utf-8", newline="")
        maps = {"row_map": csv.writer(fh["row_map"]), "review_map": csv.writer(fh["review_map"])}
        for w in maps.values():
            w.writerow(["row_no", "source_ref", "sha256"])
    tl = _TL(R.tlevel, fh.get("transformations.jsonl"))
    L = _Lines(snap, enc)
    expect, seen = 1, 0
    try:
        for u in _iter_units(R, L):
            led["units"] += 1
            if u["start"] != expect and led["coverage_ok"]:
                led["coverage_ok"] = False
                led["coverage_note"] = f"gap/overlap at line {u['start']} (expected {expect})"
            expect = u["end"] + 1
            kind = u["kind"]
            if kind == "structural":
                led["structural"][u["reason"]] += 1
                if fh and u["reason"] != "blank_line":
                    fh["structural.jsonl"].write(json.dumps(
                        {"lines": [u["start"], u["end"]], "reason": u["reason"], "raw": u["raw"]},
                        ensure_ascii=True) + "\n")
            else:
                ref = f"L{u['start']}" if u["start"] == u["end"] else f"L{u['start']}-{u['end']}"
                cells = None
                if kind == "data":
                    cells, reasons = _build_row(u, R, ref, tl)
                    if reasons:
                        kind, u["reason"] = "review", "; ".join(reasons)[:500]
                if kind == "data":
                    led["data"] += 1
                    enc_cells = [enc_cell(v) for v in cells]
                    if len(led["sample_rows"]) < 4:
                        led["sample_rows"].append((ref, cells))
                    if fh:
                        fh["expected_rows.jsonl"].write(json.dumps({"r": ref, "c": enc_cells}, ensure_ascii=True) + "\n")
                        maps["row_map"].writerow([led["data"], ref, _row_hash(ref, enc_cells)])
                else:
                    led["review"] += 1
                    key = re.sub(r"[\"'].*?[\"']|\d+", "#", u["reason"].split(";")[0])[:70]
                    led["review_reasons"][key] += 1
                    led["examples"].setdefault(key, (ref, u["reason"][:140], u["raw"][:120]))
                    shown = _sheet_text(u["raw"])
                    reason = u["reason"] + (" (raw_text escaped/truncated)" if shown != u["raw"] else "")
                    if fh:
                        enc_cells = [enc_cell(x) for x in (ref, u["start"], u["end"], reason, shown)]
                        fh["review_rows.jsonl"].write(json.dumps({"r": ref, "c": enc_cells}, ensure_ascii=True) + "\n")
                        fh["review_raw.jsonl"].write(json.dumps(
                            {"ref": ref, "lines": [u["start"], u["end"]], "reason": u["reason"], "raw": u["raw"]},
                            ensure_ascii=True) + "\n")
                        maps["review_map"].writerow([led["review"], ref, _row_hash(ref, enc_cells)])
            seen += 1
            if limit and seen >= limit:
                break
        else:
            led["complete"] = True
        if led["complete"] and expect - 1 != L.last and led["coverage_ok"]:
            led["coverage_ok"] = False
            led["coverage_note"] = f"units end at line {expect - 1} but file has {L.last} lines"
        led["lines"] = expect - 1
    finally:
        L.close()
        for f in fh.values():
            f.close()
    led["trans"] = tl.counts
    return led


# =========================================================== inspect
def _fixed_slices(lines):
    width = max(len(l) for l in lines)
    gap = [all(i >= len(l) or l[i] == " " for l in lines) for i in range(width)]
    starts = [i for i in range(width) if not gap[i] and (i == 0 or gap[i - 1])]
    return [[s, (starts[k + 1] if k + 1 < len(starts) else None)] for k, s in enumerate(starts)]


def _looks_num(s):
    return bool(re.fullmatch(r"[+-]?[\d.,]+%?", s.strip()))


def inspect_file(path):
    src = resolve_source(path)
    snap, sha = take_snapshot(src)
    size = os.path.getsize(snap)
    enc, why = detect_encoding(snap)
    eol, total, blank, nonascii, ctrl, maxlen, sumlen, nb = Counter(), 0, 0, 0, 0, 0, 0, 0
    first, last, res, rnd, sample = [], deque(maxlen=5), [], random.Random(0), []
    with open(snap, "r", encoding=enc, newline="") as f:
        for n, raw in enumerate(f, 1):
            total = n
            if raw.endswith("\r\n"):
                eol["CRLF"] += 1
            elif raw.endswith("\n"):
                eol["LF"] += 1
            elif raw.endswith("\r"):
                eol["CR"] += 1
            else:
                eol["no-terminator (last line)"] += 1
            text = _eol(raw)
            if not text.strip():
                blank += 1
                continue
            nb += 1
            maxlen, sumlen = max(maxlen, len(text)), sumlen + len(text)
            nonascii += 0 if text.isascii() else 1
            ctrl += 1 if _ILLEGAL.search(text) else 0
            item = (n, text)
            if len(first) < 10:
                first.append(item)
            else:
                last.append(item)
                if len(res) < 8:
                    res.append(item)
                else:
                    j = rnd.randrange(nb - 10)
                    if j < 8:
                        res[j] = item
            if len(sample) < 400:
                sample.append(text)
    out = [f"file: {os.path.basename(src)} · {size:,} bytes · sha256 {sha}",
           f"snapshot: read-only copy saved in .excel_snapshots/{sha[:16]}/ (all parsing uses the snapshot)",
           f"encoding: {enc} ({why}) · line endings: {dict(eol)}",
           f"lines: {total:,} (blank {blank:,}, non-blank {nb:,}) · avg length {sumlen // max(nb, 1)}, max {maxlen}"
           f" · lines with non-ASCII: {nonascii} · lines with control chars Excel cannot store: {ctrl}"]
    if not sample:
        return "\n".join(out + ["file has no non-blank lines"])
    stats = []
    for d, name in ((",", "comma"), ("\t", "tab"), (";", "semicolon"), ("|", "pipe")):
        counts = []
        for l in sample:
            try:
                counts.append(len(next(csv.reader([l], delimiter=d))))
            except csv.Error:
                counts.append(-1)
        modal, freq = Counter(counts).most_common(1)[0]
        stats.append((freq / len(counts), modal, d, name))
    out.append("delimiters (sample): " + " | ".join(f"{n} → {m} fields in {fr:.0%}" for fr, m, d, n in stats))
    best = max((x for x in stats if x[1] >= 2), default=max(stats))
    json_ok = sum(1 for l in sample[:50] if l.lstrip().startswith("{"))
    jsonl = json_ok >= 0.9 * len(sample[:50])
    kv = re.compile(r"^\s*[^:\s][^:]{0,60}:\s?.*$")
    kv_frac = sum(1 for l in sample if kv.match(l)) / len(sample)
    draft, guess = {"encoding": enc, "skip_blank": True}, "unknown (free-form): propose a regex"
    if jsonl:
        keys = []
        for l in sample[:50]:
            try:
                for k in json.loads(l):
                    if k not in keys:
                        keys.append(k)
            except ValueError:
                pass
        guess = "jsonl"
        draft.update(format="jsonl", columns=[{"name": k, "source": k, "type": "text"} for k in keys])
    elif best[0] >= 0.5 and best[1] >= 2:
        fr, modal, d, name = best
        rows = [r for r in (next(csv.reader([l], delimiter=d)) for l in sample[:40]) if len(r) == modal][:1]
        has_h = (len(rows[0]) == len(set(rows[0])) and all(x.strip() for x in rows[0])
                 and not any(_looks_num(x) for x in rows[0]))
        names = rows[0] if has_h else [f"col{i + 1}" for i in range(modal)]
        guess = (f"delimited ({name}, {modal} fields in {fr:.0%} of sample lines"
                 f"{'; MESSY: the other lines will go to review' if fr < 0.8 else ''}); header line: {'yes' if has_h else 'no'}")
        if any(l.startswith("#") for l in sample[:5]):
            draft["skip_regex"] = ["^#"]
        draft.update(format="delimited", delimiter=d, quotechar='"', header_lines=1 if has_h else 0,
                     columns=[{"name": (names[i] if i < len(names) else f"col{i + 1}").strip() or f"col{i + 1}",
                               "source": i, "type": "text"} for i in range(modal)])
    elif kv_frac >= 0.8 and blank > 0:
        keys = []
        for l in sample:
            m = kv.match(l)
            if m:
                k = l.split(":", 1)[0].strip()
                if k not in keys:
                    keys.append(k)
        guess = "blocks (Key: value lines separated by blank lines)"
        draft.update(format="blocks", columns=[{"name": k, "source": k, "type": "text"} for k in keys[:40]])
    else:
        sl = _fixed_slices(sample[:100])
        if 2 <= len(sl) <= 40 and "\t" not in "".join(sample[:100]):
            guess = f"fixed_width? ({len(sl)} column slices found; check the sample lines)"
            draft.update(format="fixed_width", fixed_width=sl,
                         columns=[{"name": f"col{i + 1}", "source": i, "type": "text"} for i in range(len(sl))])
    out.append(f"guess: {guess}")
    shown = sorted({n: t for n, t in first + res + list(last)}.items())
    out.append("sample lines (repr):")
    out += [f"  L{n}: {repr(t)[:150]}" for n, t in shown]
    out.append("draft rules (all columns typed text; refine types/cleaning, see excel_rules_help):")
    out.append(json.dumps(draft, ensure_ascii=False, separators=(",", ":"))[:2200])
    out.append("next: excel_validate_rules(path, rules). Nothing has been changed or written yet.")
    return "\n".join(out)


# =========================================================== validate (dry run)
def validate(path, rules_in, max_units=2000):
    src = resolve_source(path)
    R = Rules(rules_in)
    snap, sha = take_snapshot(src)
    enc = R.enc or detect_encoding(snap)[0]
    try:
        led = run_pipeline(R, snap, enc, max_units)
    except UnicodeDecodeError as e:
        return f"ERROR: file is not valid {enc}: {e}. Choose another encoding in rules."
    _validated[(R.hash, sha)] = {"complete": led["complete"]}
    st = led["structural"]
    total = led["data"] + led["review"]
    pct = 100 * led["review"] / total if total else 0
    out = [f"rules OK · format {R.fmt} · encoding {enc} · {len(R.cols)} columns",
           f"checked: {'whole file' if led['complete'] else f'first {led['units']} units (use max_units=0 for the whole file)'}"
           f" · lines covered: 1-{led['lines']}",
           f"data {led['data']} · review {led['review']} ({pct:.1f}%) · structural {sum(st.values())} "
           f"({', '.join(f'{k} {v}' for k, v in st.items()) or 'none'})"]
    if not led["coverage_ok"]:
        out.append(f"!! line coverage problem: {led['coverage_note']}")
    if led["review_reasons"]:
        out.append("review reasons (these records are kept, not exported as data):")
        for key, n in led["review_reasons"].most_common(8):
            ref, reason, raw = led["examples"][key]
            out.append(f"  {n}× {key}\n     e.g. {ref}: {reason}\n     raw: {raw!r}")
        out.append("→ if these are parse problems, fix the rules; if the data is genuinely ambiguous, ask the user.")
    if led["sample_rows"]:
        out.append("first parsed rows:")
        for ref, cells in led["sample_rows"]:
            out.append(f"  {ref}: " + " | ".join(f"{c['name']}={_short(v, 30)}" for c, v in zip(R.cols, cells)))
    if led["trans"]:
        out.append("transformations so far: " + ", ".join(f"{k} {v}" for k, v in sorted(led["trans"].items())))
    out.append("next: excel_build(path, rules) when the review rate is acceptable.")
    return "\n".join(out)


# =========================================================== write workbook
class _Series:
    """One logical table that rolls over to a new sheet at Excel's row limit."""

    def __init__(self, wb, base, header, types):
        self.wb, self.base, self.header, self.types = wb, base, header, types
        self.sheets, self.counts, self.ws, self.rows, self.total = [], [], None, 0, 0
        self._new()

    def _new(self):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
        title = self.base if not self.sheets else f"{self.base[:27]}_{len(self.sheets) + 1}"
        ws = self.wb.create_sheet(title=title)
        ws.freeze_panes = "A2"
        for i, h in enumerate(self.header, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(h) + 3, 12), 40)
        cells = []
        for h in self.header:
            c = WriteOnlyCell(ws, value=h)
            c.font = Font(bold=True)
            cells.append(c)
        ws.append(cells)
        self.ws, self.rows = ws, 1
        self.sheets.append(ws.title)
        self.counts.append(0)

    def add(self, values):
        from openpyxl.cell import WriteOnlyCell
        if self.rows >= MAX_ROWS:
            self._new()
        row = []
        for v in values:
            if isinstance(v, str) and v.startswith("="):     # keep text that looks like a formula as text
                c = WriteOnlyCell(self.ws, value=v)
                c.data_type = "s"
                row.append(c)
            else:
                row.append(v)
        self.ws.append(row)
        self.rows += 1
        self.total += 1
        self.counts[-1] += 1


def _jsonl(path):
    with open(path, encoding="utf-8") as f:
        for line in f:
            yield json.loads(line)


def write_workbook(R, out_path, audit, led, man):
    from openpyxl import Workbook
    wb = Workbook(write_only=True)
    header = (["source_ref"] if R.include_ref else []) + [c["name"] for c in R.cols]
    types = (["text"] if R.include_ref else []) + [c["type"] for c in R.cols]
    data = _Series(wb, R.sheet, header, types)
    for e in _jsonl(os.path.join(audit, "expected_rows.jsonl")):
        vals = [dec_cell(x) for x in e["c"]]
        data.add(([e["r"]] if R.include_ref else []) + vals)
    review = None
    if led["review"]:
        review = _Series(wb, "Review", REVIEW_HEADER, REVIEW_TYPES)
        for e in _jsonl(os.path.join(audit, "review_rows.jsonl")):
            review.add([dec_cell(x) for x in e["c"]])
    for s in [data] + ([review] if review else []):
        for ws, n in zip([wb[nm] for nm in s.sheets], s.counts):
            ws.auto_filter.ref = f"A1:{_col_letter(len(s.header))}{n + 1}"
    summ = wb.create_sheet("Summary")
    summ.column_dimensions["A"].width = 34
    summ.column_dimensions["B"].width = 90
    for k, v in man["summary_rows"]:
        summ.append([k, v])
    wb.save(out_path)
    man["data_sheets"], man["review_sheets"] = data.sheets, (review.sheets if review else [])
    man["header"], man["types"] = header, types
    return data.total, (review.total if review else 0)


def _col_letter(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


# =========================================================== verify (read back what was saved)
def _verify_series(wb, names, header, types, audit, exp_file, map_file, ref_in_row, label, probs, flags):
    ncols, n = len(header), 0
    miss = [nm for nm in names if nm not in wb.sheetnames]
    if miss:
        probs.append(f"L1 {label}: sheets missing in workbook: {miss}")
        flags["l1"] = False
        return 0
    with open(os.path.join(audit, exp_file), encoding="utf-8") as ef, \
            open(os.path.join(audit, map_file), newline="", encoding="utf-8") as mf:
        exp, mp = iter(ef), csv.reader(mf)
        next(mp, None)
        for nm in names:
            rows = wb[nm].iter_rows(values_only=True)
            hdr = list(next(rows, None) or [])
            if hdr[:ncols] != header or any(x is not None for x in hdr[ncols:]):
                probs.append(f"L1 {label} sheet {nm}: header differs: {hdr[:ncols]} vs {header}")
                flags["l1"] = False
            for row in rows:
                n += 1
                line, mrow = next(exp, None), next(mp, None)
                if line is None:
                    flags["l1"] = False
                    if len(probs) < 20:
                        probs.append(f"L1 {label}: extra row {n} in workbook that was not expected")
                    continue
                e = json.loads(line)
                vals = list(row)
                extra = [x for x in vals[ncols:] if x is not None]
                vals = (vals + [None] * ncols)[:ncols]
                act = [enc_cell(_back(v, t)) for v, t in zip(vals, types)]
                if extra:
                    flags["l1"] = False
                    probs.append(f"L1 {label} row {n}: unexpected extra cells {extra[:3]}")
                if ref_in_row:
                    exp_cells, act_cells = [["s", e["r"]]] + e["c"], act[1:]
                    if vals[0] != e["r"]:
                        flags["l1"] = False
                        if len(probs) < 20:
                            probs.append(f"L1 {label} row {n}: source_ref {vals[0]!r} != expected {e['r']!r} (order/identity)")
                else:
                    exp_cells, act_cells = e["c"], act
                if act != exp_cells:
                    flags["l2"] = False
                    for j, (a, x) in enumerate(zip(act, exp_cells)):
                        if a != x and len(probs) < 20:
                            probs.append(f"L2 {label} row {n} ({e['r']}) col {header[j]!r}: expected {x} got {a}")
                h = _row_hash(e["r"], act_cells)
                if mrow is None or h != mrow[2]:
                    flags["l3"] = False
                    if len(probs) < 20:
                        probs.append(f"L3 {label} row {n} ({e['r']}): fingerprint differs")
        left = sum(1 for _ in exp)
        if left:
            flags["l1"] = False
            probs.append(f"L1 {label}: {left} expected rows are missing from the workbook")
    return n


def verify_workbook(out_path, audit, man):
    from openpyxl import load_workbook
    flags, probs = {"l1": True, "l2": True, "l3": True}, []
    wb = load_workbook(out_path, read_only=True, data_only=True)
    try:
        exp_sheets = man["data_sheets"] + man["review_sheets"] + ["Summary"]
        if wb.sheetnames != exp_sheets:
            flags["l1"] = False
            probs.append(f"L1 sheet names {wb.sheetnames} != expected {exp_sheets}")
        nd = _verify_series(wb, man["data_sheets"], man["header"], man["types"], audit,
                            "expected_rows.jsonl", "row_map.csv", man["include_ref"], "Data", probs, flags)
        nr = 0
        if man["review_sheets"]:
            nr = _verify_series(wb, man["review_sheets"], REVIEW_HEADER, REVIEW_TYPES, audit,
                                "review_rows.jsonl", "review_map.csv", False, "Review", probs, flags)
    finally:
        wb.close()
    if nd != man["counts"]["data"] or nr != man["counts"]["review"]:
        flags["l1"] = False
        probs.append(f"L1 row counts: workbook has data {nd}/review {nr}, ledger says "
                     f"{man['counts']['data']}/{man['counts']['review']}")
    return {"level1": flags["l1"], "level2": flags["l2"], "level3": flags["l3"],
            "ok": all(flags.values()), "data_rows": nd, "review_rows": nr, "problems": probs}


# =========================================================== reports
def _verdict(v):
    return "PASS" if v else "FAIL"


def _qa_report(man, ver, led, src_unchanged):
    c = man["counts"]
    lines = [
        f"# QA report: {man['output']}",
        f"status: {'VERIFIED' if ver['ok'] else 'FAILED VERIFICATION'}",
        f"source: {man['source']}", f"source sha256: {man['source_sha256']}",
        f"source unchanged since snapshot: {src_unchanged}", f"rules sha256: {man['rules_sha256']}",
        "", "## Reconciliation ledger",
        f"source units discovered: {c['units']}  (source lines: {c['lines']})",
        f"exported to Data sheet:  {c['data']}", f"sent to Review sheet:    {c['review']}",
        f"structural (explicit):   {c['structural']}  {man['structural_by_reason']}",
        f"invariant units = exported + review + structural: {_verdict(c['units'] == c['data'] + c['review'] + c['structural'])}",
        f"line coverage (every line in exactly one unit): {_verdict(led['coverage_ok'])}",
        "", "## Read-back verification",
        f"level 1 structure (sheets, headers, counts, source ids in order): {_verdict(ver['level1'])}",
        f"level 2 exact cell values and types: {_verdict(ver['level2'])}",
        f"level 3 row fingerprints (sha256, type aware): {_verdict(ver['level3'])}",
        f"rows read back: data {ver['data_rows']}, review {ver['review_rows']}",
    ]
    if ver["problems"]:
        lines += ["", "## Problems"] + [f"- {p}" for p in ver["problems"]]
    lines += ["", "## Transformations (counts)"] + [f"- {k}: {v}" for k, v in sorted(man["transformations"].items())] \
        + ["", "## Review reasons"] + [f"- {n}× {k}" for k, n in man["review_reasons"].items()] \
        + ["", "## Notes",
           "- Excel cannot store an empty string distinct from a blank cell: empty values are blank cells (counted as empty_to_blank).",
           "- Line endings inside cell text are normalised to \\n (newline_normalized).",
           "- Exact original text of every review record: review_raw.jsonl. Non-blank structural lines: structural.jsonl."]
    return "\n".join(lines)


# =========================================================== build (orchestrator)
def _manifest_summary(man):
    c = man["counts"]
    return [("Source file", man["source"]), ("Source SHA-256", man["source_sha256"]),
            ("Encoding", man["encoding"]), ("Rules SHA-256", man["rules_sha256"]),
            ("Source lines", c["lines"]), ("Source units", c["units"]), ("Exported (Data)", c["data"]),
            ("Review queue (Review)", c["review"]), ("Structural (not data)", c["structural"]),
            ("Structural by reason", json.dumps(man["structural_by_reason"])),
            ("Transformations", json.dumps(man["transformations"])),
            ("Verification", "read-back verification result is in the QA report file (qa_report.md)"),
            ("Created", man["created"])]


def build(path, rules_in, output_path="", max_review_percent=5.0):
    t0 = time.time()
    src = resolve_source(path)
    R = Rules(rules_in)
    snap, sha = take_snapshot(src)
    if (R.hash, sha) not in _validated:
        return ("ERROR: these exact rules have not been validated on this file yet. "
                "Run excel_validate_rules(path, rules) first, review its output, then build.")
    enc = R.enc or detect_encoding(snap)[0]
    out = resolve_output(output_path, src)
    audit = os.path.splitext(out)[0] + "_audit"
    os.makedirs(audit, exist_ok=True)
    try:
        led = run_pipeline(R, snap, enc, 0, audit)
    except UnicodeDecodeError as e:
        return f"ERROR: file is not valid {enc}: {e}. Choose another encoding in rules."
    nstruct = sum(led["structural"].values())
    inv = led["units"] == led["data"] + led["review"] + nstruct
    if not (inv and led["coverage_ok"]):
        return ("FAILED before writing: reconciliation ledger does not balance "
                f"(units {led['units']} vs {led['data']}+{led['review']}+{nstruct}; coverage: "
                f"{led['coverage_note'] or 'ok'}). This is a parser bug; nothing was exported.")
    total = led["data"] + led["review"]
    pct = 100 * led["review"] / total if total else 0
    if pct > max_review_percent:
        lines = [f"STOPPED before writing: {led['review']} of {total} records ({pct:.1f}%) need review, above the "
                 f"{max_review_percent}% limit. Top reasons:"]
        for key, n in led["review_reasons"].most_common(6):
            ref, reason, raw = led["examples"][key]
            lines.append(f"  {n}× {key}  e.g. {ref}: {raw!r}")
        lines.append("Fix the rules, or ask the user whether to accept the review queue "
                     "(then call excel_build again with a higher max_review_percent).")
        return "\n".join(lines)

    man = {"source": src, "source_sha256": sha, "snapshot": snap, "encoding": enc,
           "rules": R.d, "rules_sha256": R.hash, "output": os.path.basename(out),
           "include_ref": R.include_ref, "created": time.strftime("%Y-%m-%d %H:%M:%S"),
           "counts": {"lines": led["lines"], "units": led["units"], "data": led["data"],
                      "review": led["review"], "structural": nstruct},
           "structural_by_reason": dict(led["structural"]), "transformations": dict(led["trans"]),
           "review_reasons": dict(led["review_reasons"])}
    man["summary_rows"] = _manifest_summary(man)
    with open(os.path.join(audit, "source.sha256"), "w") as f:
        f.write(f"{sha}  {os.path.basename(src)}\n")
    partial = out[:-5] + ".partial.xlsx"
    ver = None
    for attempt in (1, 2):
        write_workbook(R, partial, audit, led, man)
        ver = verify_workbook(partial, audit, man)
        if ver["ok"]:
            break
    final = out
    if ver["ok"]:
        os.replace(partial, out)
    else:
        final = out[:-5] + ".UNVERIFIED.xlsx"
        os.replace(partial, final)
    src_unchanged = sha256_file(src) == sha
    man["verification"] = {k: ver[k] for k in ("ok", "level1", "level2", "level3", "data_rows", "review_rows")}
    man["source_unchanged"] = src_unchanged
    sr = man.pop("summary_rows")
    with open(os.path.join(audit, "manifest.json"), "w", encoding="utf-8") as f:
        json.dump(man, f, indent=1, ensure_ascii=False)
    with open(os.path.join(audit, "rules.json"), "w", encoding="utf-8") as f:
        json.dump(R.d, f, indent=1, ensure_ascii=False)
    with open(os.path.join(audit, "qa_report.md"), "w", encoding="utf-8") as f:
        f.write(_qa_report(man, ver, led, src_unchanged))
    c = man["counts"]
    rel = lambda p: os.path.relpath(p, _root()).replace(os.sep, "/")
    out_lines = [
        f"{'VERIFIED' if ver['ok'] else 'FAILED VERIFICATION'} · {rel(final)} · {time.time() - t0:.1f}s",
        f"ledger: {c['units']} source units = {c['data']} exported + {c['review']} review + {c['structural']} structural "
        f"({', '.join(f'{k} {v}' for k, v in led['structural'].items()) or 'none'}) · line coverage {_verdict(led['coverage_ok'])}",
        f"read-back: structure {_verdict(ver['level1'])} · cell values/types {_verdict(ver['level2'])} · "
        f"row fingerprints {_verdict(ver['level3'])} ({ver['data_rows']} data + {ver['review_rows']} review rows checked)",
        f"source sha256 {sha[:16]}… · source unchanged: {src_unchanged}",
        f"transformations: {', '.join(f'{k} {v}' for k, v in sorted(led['trans'].items())) or 'none'}",
        f"audit folder: {rel(audit)}/ (qa_report.md, manifest.json, source.sha256, rules.json, row_map.csv, "
        "expected_rows.jsonl, review_raw.jsonl, structural.jsonl, transformations.jsonl)"]
    if led["review"]:
        out_lines.append(f"review queue: {led['review']} records on the 'Review' sheet with raw text; reasons: "
                         + "; ".join(f"{n}× {k}" for k, n in led["review_reasons"].most_common(4)))
    if not ver["ok"]:
        out_lines.append("DO NOT report success. Problems:\n  " + "\n  ".join(ver["problems"][:10]))
    return "\n".join(out_lines)


def verify_existing(output_path):
    out = resolve_output(output_path, "x.txt") if output_path else None
    if not out or not os.path.isfile(out):
        raise FileNotFoundError(f"workbook not found: {out}")
    base = out[:-len(".UNVERIFIED.xlsx")] if out.endswith(".UNVERIFIED.xlsx") else out[:-5]
    audit = base + "_audit"
    with open(os.path.join(audit, "manifest.json"), encoding="utf-8") as f:
        man = json.load(f)
    ver = verify_workbook(out, audit, man)
    snap_ok = os.path.isfile(man["snapshot"]) and sha256_file(man["snapshot"]) == man["source_sha256"]
    src_ok = os.path.isfile(man["source"]) and sha256_file(man["source"]) == man["source_sha256"]
    lines = [f"{'VERIFIED' if ver['ok'] else 'FAILED VERIFICATION'} · {os.path.basename(out)}",
             f"structure {_verdict(ver['level1'])} · cell values/types {_verdict(ver['level2'])} · "
             f"row fingerprints {_verdict(ver['level3'])} ({ver['data_rows']} data + {ver['review_rows']} review rows)",
             f"snapshot intact: {snap_ok} · original source file unchanged: {src_ok}"]
    if ver["problems"]:
        lines.append("problems:\n  " + "\n  ".join(ver["problems"][:10]))
    return "\n".join(lines)
